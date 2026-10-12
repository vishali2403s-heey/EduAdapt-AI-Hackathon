
import os
import secrets
from collections import defaultdict
from pathlib import Path

from dotenv import load_dotenv
from flask import Flask, jsonify, render_template, request, session
from pypdf import PdfReader
from pptx import Presentation
from werkzeug.utils import secure_filename

from tutor_engine import (
    answer_question,
    build_revision_plan,
    generate_quiz,
    next_difficulty,
)

load_dotenv()

BASE_DIR = Path(__file__).resolve().parent
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

ALLOWED_EXTENSIONS = {"pdf", "pptx", "txt"}
MAX_UPLOAD_BYTES = 8 * 1024 * 1024

app = Flask(__name__)
app.secret_key = os.getenv("FLASK_SECRET_KEY") or secrets.token_hex(32)
app.config["MAX_CONTENT_LENGTH"] = MAX_UPLOAD_BYTES

# Simple in-memory demo storage. Progress resets when the server restarts.
documents = {}
topic_stats = defaultdict(lambda: {"correct": 0, "total": 0})
quiz_sessions = {}


def current_user():
    if "user_id" not in session:
        session["user_id"] = secrets.token_urlsafe(12)
    return session["user_id"]


def allowed_file(filename):
    return (
        "." in filename
        and filename.rsplit(".", 1)[1].lower() in ALLOWED_EXTENSIONS
    )


def extract_text(path, extension):
    if extension == "txt":
        return path.read_text(encoding="utf-8", errors="ignore")[:100000]

    if extension == "pdf":
        reader = PdfReader(str(path))
        return "\n".join(
            page.extract_text() or "" for page in reader.pages[:100]
        )[:100000]

    if extension == "pptx":
        presentation = Presentation(str(path))
        parts = []
        for slide in presentation.slides[:100]:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    parts.append(shape.text.strip())
        return "\n".join(parts)[:100000]

    return ""


@app.route("/")
def dashboard():
    user_id = current_user()
    user_doc = documents.get(user_id, {})
    stats = {
        topic: values.copy()
        for (uid, topic), values in topic_stats.items()
        if uid == user_id
    }
    total = sum(item["total"] for item in stats.values())
    correct = sum(item["correct"] for item in stats.values())

    return render_template(
        "dashboard.html",
        filename=user_doc.get("filename", ""),
        has_material=bool(user_doc.get("text")),
        topics=stats,
        total_answers=total,
        accuracy=round(correct / total * 100) if total else 0,
        revision_plan=build_revision_plan(stats),
    )


@app.route("/upload", methods=["POST"])
def upload_document():
    user_id = current_user()
    file = request.files.get("document")

    if not file or not file.filename:
        return jsonify(error="Choose a PDF, PPTX, or TXT file."), 400

    filename = secure_filename(file.filename)
    if not filename or not allowed_file(filename):
        return jsonify(error="Only PDF, PPTX, and TXT files are supported."), 400

    extension = filename.rsplit(".", 1)[1].lower()
    safe_name = f"{secrets.token_hex(8)}.{extension}"
    path = UPLOAD_DIR / safe_name

    try:
        file.save(path)
        text = extract_text(path, extension)
        if not text.strip():
            return jsonify(
                error="No readable text found. Scanned PDFs may need OCR."
            ), 400

        documents[user_id] = {"filename": filename, "text": text}
        return jsonify(
            message="Document uploaded successfully.",
            filename=filename,
            characters=len(text),
        )
    except Exception:
        app.logger.exception("Document processing failed")
        return jsonify(error="Could not read that file. Try another document."), 400
    finally:
        # Keep only extracted text in memory for this demo.
        if path.exists():
            path.unlink()


@app.route("/ask", methods=["POST"])
def ask_tutor():
    data = request.get_json(silent=True) or {}
    question = str(data.get("question", "")).strip()

    if not question or len(question) > 2000:
        return jsonify(error="Enter a question under 2,000 characters."), 400

    material = documents.get(current_user(), {}).get("text", "")
    answer = answer_question(question, material)
    return jsonify(answer=answer)


@app.route("/quiz", methods=["POST"])
def start_quiz():
    user_id = current_user()
    data = request.get_json(silent=True) or {}

    topic = str(data.get("topic", "General"))[:100].strip() or "General"
    difficulty = str(data.get("difficulty", "medium")).lower()
    if difficulty not in {"easy", "medium", "hard"}:
        difficulty = "medium"

    material = documents.get(user_id, {}).get("text", "")
    questions = generate_quiz(
        material=material,
        topic=topic,
        difficulty=difficulty,
        count=5,
    )

    quiz_id = secrets.token_urlsafe(16)
    quiz_sessions[quiz_id] = {
        "user_id": user_id,
        "topic": topic,
        "difficulty": difficulty,
        "questions": questions,
        "answers": [],
        "index": 0,
    }

    first = questions[0]
    return jsonify(
        quiz_id=quiz_id,
        total=len(questions),
        question={
            "id": first["id"],
            "question": first["question"],
            "options": first["options"],
            "difficulty": first["difficulty"],
        },
    )


@app.route("/quiz/answer", methods=["POST"])
def submit_quiz_answer():
    user_id = current_user()
    data = request.get_json(silent=True) or {}
    quiz_id = str(data.get("quiz_id", ""))
    quiz = quiz_sessions.get(quiz_id)

    if not quiz or quiz["user_id"] != user_id:
        return jsonify(error="Quiz not found. Start a new quiz."), 404

    index = quiz["index"]
    if index >= len(quiz["questions"]):
        return jsonify(error="This quiz has already finished."), 400

    selected = data.get("answer")
    if not isinstance(selected, int) or isinstance(selected, bool) or selected not in range(4):
        return jsonify(error="Choose one of the four options."), 400

    question = quiz["questions"][index]
    was_correct = selected == question["correct"]

    stats = topic_stats[(user_id, quiz["topic"])]
    stats["total"] += 1
    if was_correct:
        stats["correct"] += 1

    quiz["answers"].append(was_correct)
    quiz["index"] += 1
    quiz["difficulty"] = next_difficulty(
        quiz["difficulty"], was_correct
    )

    result = {
        "correct": was_correct,
        "explanation": question.get("explanation", ""),
        "next_difficulty": quiz["difficulty"],
        "finished": quiz["index"] >= len(quiz["questions"]),
    }

    if result["finished"]:
        result["score"] = sum(quiz["answers"])
        result["total"] = len(quiz["questions"])
        result["accuracy"] = round(
            result["score"] / result["total"] * 100
        )
    else:
        next_q = quiz["questions"][quiz["index"]]
        result["question"] = {
            "id": next_q["id"],
            "question": next_q["question"],
            "options": next_q["options"],
            "difficulty": next_q["difficulty"],
        }

    return jsonify(result)


@app.errorhandler(413)
def upload_too_large(_error):
    return jsonify(error="File is too large. Maximum size is 8 MB."), 413


if __name__ == "__main__":
    app.run(debug=os.getenv("FLASK_DEBUG") == "1")
